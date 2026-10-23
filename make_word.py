"""paper/main.tex -> paper/Major_Project_Paper.docx (Word version of the paper).
Pandoc cannot draw TikZ, algorithm boxes or PDF figures, so: PDF figures are rendered to PNG; the pipeline and architecture
TikZ figures and the algorithm box are compiled standalone with tectonic and rendered to PNG; \\ref/\\cite numbers are taken
from the LaTeX build's main.aux; table files are inlined without \\resizebox. Needs tectonic, pandoc, pymupdf, python-docx.
Run from the repo root after compiling the paper with --keep-intermediates --outdir build_word."""
import os, re, subprocess, pymupdf
from docx import Document
from docx.shared import Pt

ROOT = os.path.dirname(os.path.abspath(__file__)); P = os.path.join(ROOT, "paper"); B = os.path.join(ROOT, "build_word")
BS = chr(92)
os.makedirs(B, exist_ok=True)
tex = open(os.path.join(P, "main.tex"), encoding="utf-8").read()
aux = open(os.path.join(B, "main.aux"), encoding="utf-8").read()


def png_from_pdf(pdf, png, dpi=300, crop=False):
    d = pymupdf.open(pdf); pg = d[0]; clip = None
    if crop:                                                  # tight box around everything drawn on the page
        r = pymupdf.Rect()
        for b in pg.get_text("dict")["blocks"]: r |= pymupdf.Rect(b["bbox"])
        for dr in pg.get_drawings(): r |= dr["rect"]
        clip = r + (-4, -4, 4, 4)
    pg.get_pixmap(dpi=dpi, clip=clip).save(png)


def tectonic(name, body, preamble):
    src = os.path.join(B, name + ".tex")
    open(src, "w", encoding="utf-8").write(preamble + "\n" + BS + "begin{document}\n" + body + "\n" + BS + "end{document}\n")
    r = subprocess.run(["tectonic", "-X", "compile", src, "--outdir", B], cwd=P, capture_output=True, text=True)
    if r.returncode: raise RuntimeError(name + " failed:\n" + r.stderr[-2500:])
    return os.path.join(B, name + ".pdf")


# 1. figures -> PNG
for f in ("fig_offset", "fig_attacks18", "fig_families"):
    png_from_pdf(os.path.join(P, f + ".pdf"), os.path.join(B, f + ".png"))
TIKZ_PRE = (BS + "documentclass[border=6pt]{standalone}\n" + BS + "renewcommand{" + BS + "rmdefault}{lmr}\n" + BS + "usepackage{amsmath}\n"
            + BS + "usepackage{tikz}\n" + BS + "usetikzlibrary{positioning, arrows.meta, fit}")
m = re.search(re.escape(BS + "begin{tikzpicture}") + r".*?" + re.escape(BS + "end{tikzpicture}"), tex, re.S)
png_from_pdf(tectonic("w_pipeline", BS + "scriptsize\n" + m.group(0), TIKZ_PRE), os.path.join(B, "pipeline.png"))
PP = P.replace(BS, "/")
png_from_pdf(tectonic("w_arch", BS + "scriptsize\n" + BS + "input{" + PP + "/fig_arch.tex}", TIKZ_PRE), os.path.join(B, "arch.png"))
ALG_PRE = (BS + "documentclass{article}\n" + BS + "renewcommand{" + BS + "rmdefault}{lmr}\n" + BS + "usepackage[paperwidth=4in,paperheight=6in,margin=0.25in]{geometry}\n"
           + BS + "usepackage{amsmath,amssymb}\n" + BS + "usepackage{algorithm}\n" + BS + "usepackage{algpseudocode}\n" + BS + "pagestyle{empty}")
png_from_pdf(tectonic("w_alg", BS + "input{" + PP + "/alg_detect.tex}", ALG_PRE), os.path.join(B, "alg.png"), crop=True)

# 2. numbers for \ref and \cite from the LaTeX build
labels = {k: v.replace(BS + "mbox  {", "").replace("}", "").strip() for k, v in re.findall(r"\\newlabel\{([^}#]+)\}\{\{((?:\\mbox  \{)?[^}]*\}?)", aux)}
cites = dict(re.findall(r"\\bibcite\{([^}]+)\}\{([^}]+)\}", aux))
alg_caption = re.search(r"\\caption\{(.*?)\}\s*\\label\{alg:detect\}", open(os.path.join(P, "alg_detect.tex"), encoding="utf-8").read(), re.S)
alg_caption = alg_caption.group(1) if alg_caption else "Receiver-side detection and training"

# 3. rewrite the source for pandoc
w = tex
img = lambda f, width="3.4in": BS + "includegraphics[width=" + width + "]{" + os.path.join(B, f).replace(BS, "/") + "}"
w = w.replace(m.group(0), img("pipeline.png"))
w = w.replace(BS + "input{fig_arch.tex}", img("arch.png"))
w = w.replace(BS + "input{alg_detect.tex}", BS + "begin{figure}" + BS + "centering\n" + img("alg.png") + "\n" + BS + "caption{Algorithm 1: " + alg_caption + "}\n" + BS + "end{figure}")
for f in ("fig_offset", "fig_attacks18", "fig_families"):
    w = re.sub(re.escape(BS + "includegraphics[width=" + BS + "columnwidth]{" + f + ".pdf}"), lambda _: img(f + ".png"), w)
def inline_table(mt):
    t = open(os.path.join(P, mt.group(1)), encoding="utf-8").read()
    t = re.sub(r"\\resizebox\{\\columnwidth\}\{!\}\{(\\begin\{tabular\}.*?\\end\{tabular\})\}", lambda x: x.group(1), t, flags=re.S)
    return t
w = re.sub(r"\\input\{(tables_[a-z0-9]+\.tex)\}", inline_table, w)
w = re.sub(r"\\resizebox\{\\columnwidth\}\{!\}\{(\\begin\{tabular\}.*?\\end\{tabular\})\}", lambda x: x.group(1), w, flags=re.S)
def number_caption(mt):                                       # "Fig. N." / "TABLE N." in front of each caption (pandoc drops LaTeX numbering)
    env = mt.group(0); lab = re.search(r"\\label\{([^}]+)\}", env)
    if not lab or lab.group(1) not in labels: return env
    pre = ("Fig. " + labels[lab.group(1)] + ". ") if mt.group(1) == "figure" else ("TABLE " + labels[lab.group(1)] + ". ")
    return env.replace(BS + "caption{", BS + "caption{" + pre, 1)
w = re.sub(r"\\begin\{(figure|table)\}.*?\\end\{\1\}", number_caption, w, flags=re.S)
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]; sec = [0, 0]
def number_heading(mt):                                       # IEEE-style "III." and "A." headings, as the text refers to them
    if mt.group(1) == "section": sec[0] += 1; sec[1] = 0; return BS + "section*{" + ROMAN[sec[0] - 1] + ". " + mt.group(2) + "}"
    sec[1] += 1; return BS + "subsection*{" + "ABCDEFGH"[sec[1] - 1] + ". " + mt.group(2) + "}"
w = re.sub(r"\\(section|subsection)\{([^}]*)\}", number_heading, w)
w = re.sub(r"\\ref\{([^}]+)\}", lambda x: labels.get(x.group(1), "?"), w)
w = re.sub(r"\\cite\{([^}]+)\}", lambda x: "[" + ", ".join(cites.get(k.strip(), "?") for k in x.group(1).split(",")) + "]", w)
w = re.sub(r"\\author\{.*?\\maketitle", lambda _: BS + "author{Shlok (UAI05BTA2317), Jiya (UAI05BTA2313), Caleb (UAI05BTA2319) -- BTech AI ML, Universal AI University" + BS + BS + " Guides: Rajesh Bhise and Laxman Buddhewad}\n" + BS + "maketitle", w, flags=re.S)
w = re.sub(r"\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}", lambda x: BS + "noindent" + BS + "textbf{Index Terms}: " + x.group(1).strip(), w, flags=re.S)
bib = re.search(r"\\begin\{thebibliography\}\{\d+\}(.*?)\\end\{thebibliography\}", w, re.S)
items = re.split(r"\\bibitem\{([^}]+)\}", bib.group(1))[1:]
refs = "\n\n".join("[" + cites.get(items[i], "?") + "] " + items[i + 1].strip() for i in range(0, len(items), 2))
w = w.replace(bib.group(0), BS + "section*{References}\n" + refs)
for pk in ("algorithm", "algpseudocode", "tikz"):
    w = re.sub(r"\\usepackage(\[[^\]]*\])?\{" + pk + r"\}\n", "", w)
w = re.sub(r"\\usetikzlibrary\{[^}]*\}\n", "", w)
open(os.path.join(B, "main_word.tex"), "w", encoding="utf-8").write(w)

# 4. reference style (Times New Roman, IEEE-like sizes) and conversion
ref = os.path.join(B, "reference.docx")
subprocess.run(["pandoc", "-o", ref, "--print-default-data-file", "reference.docx"], check=True)
doc = Document(ref)
# body text 12 pt (= 16 px at 96 dpi, the professor's requirement); headings scaled up, captions/tables/references slightly smaller
for name, size, bold in (("Normal", 12, None), ("Body Text", 12, None), ("First Paragraph", 12, None), ("Compact", 11, None), ("Title", 18, True),
                         ("Author", 12, None), ("Abstract", 12, None), ("Heading 1", 14, True), ("Heading 2", 12, True), ("Heading 3", 12, True),
                         ("Caption", 10, None), ("Image Caption", 10, None), ("Table Caption", 10, None), ("Bibliography", 11, None)):
    try:
        st = doc.styles[name]; st.font.name = "Times New Roman"; st.font.size = Pt(size)
        rf = st.element.rPr.rFonts                            # drop theme-font links, which otherwise override the name (headings stayed sans)
        for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"): rf.attrib.pop("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + a, None)
        for a in ("ascii", "hAnsi", "cs", "eastAsia"): rf.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + a, "Times New Roman")
        if bold is not None: st.font.bold = bold
        st.font.color.rgb = None
    except KeyError:
        pass
doc.save(ref)
out = os.path.join(P, "Major_Project_Paper.docx")
r = subprocess.run(["pandoc", os.path.join(B, "main_word.tex"), "-f", "latex", "-o", out, "--reference-doc", ref, "--resource-path", P],
                   cwd=P, capture_output=True, text=True)
print(r.stderr[-3000:] if r.stderr else "pandoc ok")

# 5. IEEE-style table rules (top, below header, bottom) and centred figures, applied to the generated file
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
def border(el, side, sz):
    b = OxmlElement("w:" + side); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "0"); b.set(qn("w:color"), "000000"); el.append(b)
d = Document(out)
for t in d.tables:
    tblPr = t._tbl.tblPr; old = tblPr.find(qn("w:tblBorders"))
    if old is not None: tblPr.remove(old)
    bd = OxmlElement("w:tblBorders"); border(bd, "top", 8); border(bd, "bottom", 8); tblPr.append(bd)
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr(); cb = OxmlElement("w:tcBorders"); border(cb, "bottom", 4); tcPr.append(cb)
for p in d.paragraphs:
    if p._p.xpath(".//pic:pic"): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.save(out); print("wrote", out)
